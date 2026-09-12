"""
多模态文件解析模块（简化版）
支持PDF、Word、PPT等格式的解析
"""

import os
import logging
from typing import Dict, List, Optional
import PyPDF2
from docx import Document
from pptx import Presentation

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class DocumentParser:
    """多模态文档解析器（简化版）"""
    
    def __init__(self):
        """初始化解析器"""
        self.supported_formats = {
            'pdf': self._parse_pdf,
            'docx': self._parse_word,
            'pptx': self._parse_pptx,
            'txt': self._parse_text
        }
        logger.info("文档解析器初始化完成")
    
    def parse(self, file_path: str) -> Dict:
        """
        解析文件
        
        Args:
            file_path: 文件路径
            
        Returns:
            解析结果字典，包含文本内容、元数据等
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"文件不存在: {file_path}")
        
        # 获取文件扩展名
        file_ext = os.path.splitext(file_path)[1].lower().lstrip('.')
        
        if file_ext not in self.supported_formats:
            # 不支持的格式，返回基本信息
            return {
                'type': file_ext,
                'file_name': os.path.basename(file_path),
                'file_type': file_ext,
                'file_size': os.path.getsize(file_path),
                'full_text': f"[{file_ext}格式文件，暂不支持详细解析]"
            }
        
        # 调用对应的解析函数
        parser_func = self.supported_formats[file_ext]
        result = parser_func(file_path)
        
        # 添加基础元数据
        result['file_name'] = os.path.basename(file_path)
        result['file_type'] = file_ext
        result['file_size'] = os.path.getsize(file_path)
        
        logger.info(f"成功解析文件: {file_path}, 类型: {file_ext}")
        return result
    
    def _parse_pdf(self, file_path: str) -> Dict:
        """
        解析PDF文件
        
        Args:
            file_path: PDF文件路径
            
        Returns:
            解析结果
        """
        text_content = []
        
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                num_pages = len(pdf_reader.pages)
                
                for page_num in range(num_pages):
                    page = pdf_reader.pages[page_num]
                    text = page.extract_text()
                    if text.strip():
                        text_content.append({
                            'page': page_num + 1,
                            'content': text
                        })
            
            full_text = '\n'.join([item['content'] for item in text_content])
            
            # 检查提取的文本是否为空
            if not full_text.strip():
                logger.warning(f"PDF文件文本提取为空: {file_path}")
            
            logger.info(f"PDF解析成功: {file_path}, 总页数: {num_pages}, 提取文本长度: {len(full_text)}")
            
            return {
                'type': 'pdf',
                'total_pages': num_pages,
                'text_content': text_content,
                'full_text': full_text
            }
        except Exception as e:
            logger.error(f"PDF解析失败: {e}")
            return {
                'type': 'pdf',
                'full_text': f"PDF解析失败: {str(e)}"
            }
    
    def _parse_word(self, file_path: str) -> Dict:
        """
        解析Word文档
        
        Args:
            file_path: Word文件路径
            
        Returns:
            解析结果
        """
        try:
            doc = Document(file_path)
            paragraphs = []
            
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)
            
            # 提取表格内容
            tables_content = []
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    table_data.append(row_data)
                tables_content.append(table_data)
            
            full_text = '\n'.join(paragraphs)
            
            # 检查提取的文本是否为空
            if not full_text.strip():
                logger.warning(f"Word文档文本提取为空: {file_path}")
            
            logger.info(f"Word解析成功: {file_path}, 段落数: {len(paragraphs)}, 提取文本长度: {len(full_text)}")
            
            return {
                'type': 'docx',
                'paragraphs': paragraphs,
                'tables': tables_content,
                'full_text': full_text
            }
        except Exception as e:
            logger.error(f"Word解析失败: {e}")
            return {
                'type': 'docx',
                'full_text': f"Word解析失败: {str(e)}"
            }
    
    def _parse_pptx(self, file_path: str) -> Dict:
        """
        解析PPT文件
        
        Args:
            file_path: PPT文件路径
            
        Returns:
            解析结果
        """
        try:
            prs = Presentation(file_path)
            slides_content = []
            
            for slide_num, slide in enumerate(prs.slides):
                slide_text = []
                
                # 提取文本框内容
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text)
                
                if slide_text:
                    slides_content.append({
                        'slide': slide_num + 1,
                        'content': '\n'.join(slide_text)
                    })
            
            full_text = '\n'.join([item['content'] for item in slides_content])
            
            # 检查提取的文本是否为空
            if not full_text.strip():
                logger.warning(f"PPT文件文本提取为空: {file_path}")
            
            logger.info(f"PPT解析成功: {file_path}, 幻灯片数: {len(prs.slides)}, 提取文本长度: {len(full_text)}")
            
            return {
                'type': 'pptx',
                'total_slides': len(prs.slides),
                'slides_content': slides_content,
                'full_text': full_text
            }
        except Exception as e:
            logger.error(f"PPT解析失败: {e}")
            return {
                'type': 'pptx',
                'full_text': f"PPT解析失败: {str(e)}"
            }
    
    def _parse_text(self, file_path: str) -> Dict:
        """
        解析纯文本文件
        
        Args:
            file_path: 文本文件路径
            
        Returns:
            解析结果
        """
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
            
            return {
                'type': 'txt',
                'full_text': text
            }
        except Exception as e:
            logger.error(f"文本解析失败: {e}")
            return {
                'type': 'txt',
                'full_text': f"文本解析失败: {str(e)}"
            }
    
    def batch_parse(self, file_paths: List[str]) -> List[Dict]:
        """
        批量解析文件
        
        Args:
            file_paths: 文件路径列表
            
        Returns:
            解析结果列表
        """
        results = []
        for file_path in file_paths:
            try:
                result = self.parse(file_path)
                results.append(result)
            except Exception as e:
                logger.error(f"解析文件失败 {file_path}: {e}")
                results.append({
                    'file_path': file_path,
                    'error': str(e)
                })
        
        logger.info(f"批量解析完成，成功: {len([r for r in results if 'error' not in r])}, 失败: {len([r for r in results if 'error' in r])}")
        return results


# 全局解析器实例
_document_parser = None


def get_document_parser() -> DocumentParser:
    """
    获取全局文档解析器实例（单例模式）
    
    Returns:
        DocumentParser实例
    """
    global _document_parser
    if _document_parser is None:
        _document_parser = DocumentParser()
    return _document_parser
