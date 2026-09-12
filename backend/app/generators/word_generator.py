"""
Word教案自动生成引擎
根据融合指令自动生成完整的Word教案
"""

import os
import logging
from typing import Dict, List, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class WordGenerator:
    """Word教案生成器"""
    
    def __init__(self):
        """初始化Word生成器"""
        logger.info("Word生成器初始化完成")
    
    def generate(self, fusion_instruction: Dict, output_path: str = None) -> str:
        """
        生成Word教案
        
        Args:
            fusion_instruction: 融合后的课件生成指令
            output_path: 输出文件路径
            
        Returns:
            生成的Word文件路径
        """
        # 创建Word文档
        doc = Document()
        
        # 设置页面边距
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1.25)
            section.right_margin = Inches(1.25)
        
        # 添加教案标题
        self._add_title(doc, fusion_instruction)
        
        # 添加教学基本信息
        self._add_basic_info(doc, fusion_instruction)
        
        # 添加教学目标
        self._add_teaching_objective(doc, fusion_instruction)
        
        # 添加学情分析
        self._add_student_analysis(doc, fusion_instruction)
        
        # 添加教学重难点
        self._add_key_difficulties(doc, fusion_instruction)
        
        # 添加教学方法
        self._add_teaching_methods(doc, fusion_instruction)
        
        # 添加教学流程
        self._add_teaching_flow(doc, fusion_instruction)
        
        # 添加课堂互动设计
        self._add_classroom_activities(doc, fusion_instruction)
        
        # 添加随堂练习
        self._add_class_practice(doc, fusion_instruction)
        
        # 添加课后作业
        self._add_homework(doc, fusion_instruction)
        
        # 添加板书设计
        self._add_blackboard_design(doc, fusion_instruction)
        
        # 添加教学反思
        self._add_reflection(doc, fusion_instruction)
        
        # 确保输出目录存在
        if output_path is None:
            output_dir = os.path.join(os.path.dirname(__file__), '../../outputs')
            os.makedirs(output_dir, exist_ok=True)
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            course_title = fusion_instruction.get('course_title', '教学教案')
            # 清理文件名中的非法字符
            course_title = ''.join(c for c in course_title if c.isalnum() or c in (' ', '-', '_')).strip()
            output_path = os.path.join(output_dir, f"{course_title}_教案_{timestamp}.docx")
        
        # 保存Word文档
        doc.save(output_path)
        logger.info(f"Word教案生成成功: {output_path}")
        
        return output_path
    
    def _add_title(self, doc: Document, instruction: Dict):
        """添加教案标题"""
        title = instruction.get('course_title', '教学教案')
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        doc.add_paragraph()  # 空行
    
    def _add_basic_info(self, doc: Document, instruction: Dict):
        """添加教学基本信息"""
        p = doc.add_paragraph()
        run = p.add_run('一、教学基本信息')
        run.font.size = Pt(16)
        run.font.bold = True
        
        # 创建信息表格
        table = doc.add_table(rows=5, cols=2)
        table.style = 'Table Grid'
        
        info_data = [
            ('课程名称', instruction.get('course_title', '')),
            ('授课时长', instruction.get('duration', '45分钟')),
            ('授课对象', instruction.get('target_audience', '学生')),
            ('授课教师', '教师姓名'),
            ('授课日期', datetime.now().strftime('%Y-%m-%d'))
        ]
        
        for i, (key, value) in enumerate(info_data):
            row = table.rows[i]
            row.cells[0].text = key
            row.cells[1].text = str(value)
            
            # 设置单元格样式
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.runs[0].font.size = Pt(12)
        
        doc.add_paragraph()  # 空行
    
    def _add_teaching_objective(self, doc: Document, instruction: Dict):
        """添加教学目标"""
        p = doc.add_paragraph()
        run = p.add_run('二、教学目标')
        run.font.size = Pt(16)
        run.font.bold = True
        
        objective = instruction.get('teaching_objective', '通过本课学习，学生能够掌握相关知识点，提升能力。')
        p = doc.add_paragraph(objective)
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.runs[0].font.size = Pt(12)
        
        doc.add_paragraph()  # 空行
    
    def _add_student_analysis(self, doc: Document, instruction: Dict):
        """添加学情分析"""
        p = doc.add_paragraph()
        run = p.add_run('三、学情分析')
        run.font.size = Pt(16)
        run.font.bold = True
        
        target_audience = instruction.get('target_audience', '学生')
        analysis = f"本课程面向{target_audience}，学生已具备一定的基础知识，但对本课内容可能存在理解上的困难。需要通过生动的案例和互动活动激发学习兴趣。"
        
        p = doc.add_paragraph(analysis)
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.runs[0].font.size = Pt(12)
        
        doc.add_paragraph()  # 空行
    
    def _add_key_difficulties(self, doc: Document, instruction: Dict):
        """添加教学重难点"""
        p = doc.add_paragraph()
        run = p.add_run('四、教学重难点')
        run.font.size = Pt(16)
        run.font.bold = True
        
        # 教学重点
        p = doc.add_paragraph()
        run = p.add_run('1. 教学重点：')
        run.font.size = Pt(14)
        run.font.bold = True
        
        key_points = instruction.get('key_points', ['重点内容1', '重点内容2'])
        for point in key_points:
            p = doc.add_paragraph(f"   • {point}")
            p.paragraph_format.left_indent = Inches(0.5)
            p.runs[0].font.size = Pt(12)
        
        # 教学难点
        p = doc.add_paragraph()
        run = p.add_run('2. 教学难点：')
        run.font.size = Pt(14)
        run.font.bold = True
        
        difficult_points = instruction.get('difficult_points', ['难点内容1', '难点内容2'])
        for point in difficult_points:
            p = doc.add_paragraph(f"   • {point}")
            p.paragraph_format.left_indent = Inches(0.5)
            p.runs[0].font.size = Pt(12)
        
        doc.add_paragraph()  # 空行
    
    def _add_teaching_methods(self, doc: Document, instruction: Dict):
        """添加教学方法"""
        p = doc.add_paragraph()
        run = p.add_run('五、教学方法')
        run.font.size = Pt(16)
        run.font.bold = True
        
        methods = [
            '讲授法：系统讲解知识点，确保学生理解核心概念',
            '案例分析法：通过具体案例分析，加深学生对知识点的理解',
            '互动讨论法：组织学生进行小组讨论，培养思辨能力',
            '多媒体演示法：运用PPT、视频等多媒体手段，增强教学效果'
        ]
        
        for method in methods:
            p = doc.add_paragraph(f"   • {method}")
            p.paragraph_format.left_indent = Inches(0.5)
            p.runs[0].font.size = Pt(12)
        
        doc.add_paragraph()  # 空行
    
    def _add_teaching_flow(self, doc: Document, instruction: Dict):
        """添加教学流程"""
        p = doc.add_paragraph()
        run = p.add_run('六、教学流程')
        run.font.size = Pt(16)
        run.font.bold = True
        
        teaching_flow = instruction.get('teaching_flow', [])
        if not teaching_flow:
            teaching_flow = [
                {'stage': '导入', 'content': '课程导入，激发兴趣', 'time_allocation': '5分钟'},
                {'stage': '新授', 'content': '新知识讲解', 'time_allocation': '25分钟'},
                {'stage': '练习', 'content': '课堂练习', 'time_allocation': '10分钟'},
                {'stage': '总结', 'content': '课程总结', 'time_allocation': '5分钟'}
            ]
        
        for i, flow in enumerate(teaching_flow, 1):
            stage = flow.get('stage', f'阶段{i}')
            content = flow.get('content', '')
            time = flow.get('time_allocation', '')
            
            p = doc.add_paragraph()
            run = p.add_run(f'{i}. {stage}（{time}）')
            run.font.size = Pt(14)
            run.font.bold = True
            
            p = doc.add_paragraph(f'   {content}')
            p.paragraph_format.left_indent = Inches(0.5)
            p.runs[0].font.size = Pt(12)
        
        doc.add_paragraph()  # 空行
    
    def _add_classroom_activities(self, doc: Document, instruction: Dict):
        """添加课堂互动设计"""
        p = doc.add_paragraph()
        run = p.add_run('七、课堂互动设计')
        run.font.size = Pt(16)
        run.font.bold = True
        
        activities = instruction.get('classroom_activities', [])
        # 兼容两种数据结构：字符串列表或对象列表
        if activities and isinstance(activities[0], dict):
            for i, activity in enumerate(activities, 1):
                activity_name = activity.get('activity', f'活动{i}')
                description = activity.get('description', '')
                p = doc.add_paragraph(f'   活动{i}：{activity_name}')
                p.paragraph_format.left_indent = Inches(0.5)
                p.runs[0].font.size = Pt(12)
                if description:
                    p = doc.add_paragraph(f'      {description}')
                    p.paragraph_format.left_indent = Inches(0.8)
                    p.runs[0].font.size = Pt(11)
        else:
            for i, activity in enumerate(activities, 1):
                p = doc.add_paragraph(f'   活动{i}：{activity}')
                p.paragraph_format.left_indent = Inches(0.5)
                p.runs[0].font.size = Pt(12)
        
        doc.add_paragraph()  # 空行
    
    def _add_class_practice(self, doc: Document, instruction: Dict):
        """添加随堂练习"""
        p = doc.add_paragraph()
        run = p.add_run('八、随堂练习')
        run.font.size = Pt(16)
        run.font.bold = True
        
        practice = "1. 基础练习题：巩固本课所学知识点\n2. 拓展思考题：引导学生深入思考\n3. 小组合作题：培养团队协作能力"
        
        p = doc.add_paragraph(practice)
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.runs[0].font.size = Pt(12)
        
        doc.add_paragraph()  # 空行
    
    def _add_homework(self, doc: Document, instruction: Dict):
        """添加课后作业"""
        p = doc.add_paragraph()
        run = p.add_run('九、课后作业')
        run.font.size = Pt(16)
        run.font.bold = True
        
        homework = instruction.get('homework', '完成课后练习题，预习下节课内容')
        
        p = doc.add_paragraph(homework)
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.runs[0].font.size = Pt(12)
        
        doc.add_paragraph()  # 空行
    
    def _add_blackboard_design(self, doc: Document, instruction: Dict):
        """添加板书设计"""
        p = doc.add_paragraph()
        run = p.add_run('十、板书设计')
        run.font.size = Pt(16)
        run.font.bold = True
        
        blackboard = instruction.get('blackboard_design', '板书设计待完善')
        
        p = doc.add_paragraph(blackboard)
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.runs[0].font.size = Pt(12)
        
        doc.add_paragraph()  # 空行
    
    def _add_reflection(self, doc: Document, instruction: Dict):
        """添加教学反思"""
        p = doc.add_paragraph()
        run = p.add_run('十一、教学反思')
        run.font.size = Pt(16)
        run.font.bold = True
        
        reflection = "课后根据学生反馈和教学效果，对本节课进行反思总结，为后续教学改进提供参考。"
        
        p = doc.add_paragraph(reflection)
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.runs[0].font.size = Pt(12)
        
        doc.add_paragraph()  # 空行


# 全局Word生成器实例
_word_generator = None


def get_word_generator() -> WordGenerator:
    """
    获取全局Word生成器实例（单例模式）
    
    Returns:
        WordGenerator实例
    """
    global _word_generator
    if _word_generator is None:
        _word_generator = WordGenerator()
    return _word_generator
