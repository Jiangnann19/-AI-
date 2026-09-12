"""
PPTX课件自动生成引擎
根据融合指令自动生成完整的PPTX课件
"""

import os
import logging
from typing import Dict, List, Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from datetime import datetime

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class PPTGenerator:
    """PPT课件生成器"""
    
    def __init__(self):
        """初始化PPT生成器"""
        self.slide_templates = {
            'cover': self._create_cover_slide,
            'directory': self._create_directory_slide,
            'knowledge': self._create_knowledge_slide,
            'case': self._create_case_slide,
            'interaction': self._create_interaction_slide,
            'summary': self._create_summary_slide,
            'ending': self._create_ending_slide,
            'tcp_handshake': self._create_tcp_handshake_slide
        }
        logger.info("PPT生成器初始化完成")
    
    def generate(self, fusion_instruction: Dict, output_path: str = None) -> str:
        """
        生成PPT课件
        
        Args:
            fusion_instruction: 融合后的课件生成指令
            output_path: 输出文件路径
            
        Returns:
            生成的PPT文件路径
        """
        # 检测用户是否要求只要1页
        course_title = fusion_instruction.get('course_title', '')
        special_requirements = fusion_instruction.get('special_requirements', '')
        single_page_mode = '只要1页' in special_requirements or '只要1页' in course_title or '只要一页' in special_requirements
        
        # 创建PPT对象
        prs = Presentation()
        
        # 设置幻灯片尺寸（16:9）
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(5.625)
        
        # 如果是单页模式且是TCP相关课程，直接生成单页完整课件
        if single_page_mode and ('TCP' in course_title or 'tcp' in course_title.lower()):
            logger.info("检测到单页TCP课件要求，生成单页完整教学幻灯片")
            # 为前端预览设置template_type
            if 'ppt_structure' not in fusion_instruction:
                fusion_instruction['ppt_structure'] = []
            fusion_instruction['ppt_structure'] = [{
                'slide_type': 'tcp_handshake',
                'template_type': 'tcp_handshake',
                'title': course_title,
                'content': fusion_instruction.get('knowledge_points', ['TCP三次握手'])
            }]
            self._create_tcp_handshake_single_slide(prs, fusion_instruction)
            # 单页模式不继续处理，直接跳到保存阶段
        else:
            # 获取PPT结构
            ppt_structure = fusion_instruction.get('ppt_structure', [])
            
            if not ppt_structure:
                logger.warning("未提供PPT结构，使用默认结构")
                ppt_structure = self._get_default_structure(fusion_instruction)
            
            # 根据结构生成幻灯片
            for slide_info in ppt_structure:
                slide_type = slide_info.get('slide_type', 'knowledge')
                slide_title = slide_info.get('title', '')
                
                # 检测是否是TCP相关内容，强制使用tcp_handshake模板
                if 'TCP' in slide_title or 'tcp' in slide_title.lower() or '握手' in slide_title:
                    logger.info(f"检测到TCP相关幻灯片: {slide_title}, 强制使用tcp_handshake模板")
                    slide_type = 'tcp_handshake'
                    slide_info['slide_type'] = 'tcp_handshake'
                    slide_info['template_type'] = 'tcp_handshake'
                
                template_func = self.slide_templates.get(slide_type, self._create_knowledge_slide)
                logger.info(f"使用模板: {slide_type}, 幻灯片标题: {slide_title}")
                template_func(prs, slide_info, fusion_instruction)
        
        # 确保输出目录存在
        if output_path is None:
            output_dir = os.path.join(os.path.dirname(__file__), '../../outputs')
            os.makedirs(output_dir, exist_ok=True)
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            course_title = fusion_instruction.get('course_title', '教学课件')
            # 清理文件名中的非法字符
            course_title = ''.join(c for c in course_title if c.isalnum() or c in (' ', '-', '_')).strip()
            output_path = os.path.join(output_dir, f"{course_title}_{timestamp}.pptx")
        
        # 保存PPT
        prs.save(output_path)
        logger.info(f"PPT课件生成成功: {output_path}")
        
        return output_path
    
    def _create_cover_slide(self, prs: Presentation, slide_info: Dict, instruction: Dict):
        """
        创建封面幻灯片
        
        Args:
            prs: Presentation对象
            slide_info: 幻灯片信息
            instruction: 完整指令
        """
        slide_layout = prs.slide_layouts[6]  # 空白布局
        slide = prs.slides.add_slide(slide_layout)
        
        # 添加标题
        title = slide.shapes.title
        if title is not None:
            title.text = slide_info.get('title', instruction.get('course_title', '教学课件'))
            title.text_frame.paragraphs[0].font.size = Pt(44)
            title.text_frame.paragraphs[0].font.bold = True
            title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        else:
            # 没有标题占位符，手动添加文本框作为标题
            left = Inches(0.5)
            top = Inches(0.5)
            width = Inches(9)
            height = Inches(1.2)
            txBox = slide.shapes.add_textbox(left, top, width, height)
            tf = txBox.text_frame
            tf.text = slide_info.get('title', instruction.get('course_title', '教学课件'))
            tf.paragraphs[0].font.size = Pt(44)
            tf.paragraphs[0].font.bold = True
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        # 添加副标题信息
        left = Inches(1)
        top = Inches(3)
        width = Inches(8)
        height = Inches(0.5)
        
        info_box = slide.shapes.add_textbox(left, top, width, height)
        tf = info_box.text_frame
        if tf is not None:
            tf.text = f"授课教师：教师姓名  |  授课时长：{instruction.get('duration', '45分钟')}  |  受众：{instruction.get('target_audience', '学生')}"
            if tf.paragraphs and len(tf.paragraphs) > 0:
                tf.paragraphs[0].font.size = Pt(18)
                tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        # 添加日期
        date_box = slide.shapes.add_textbox(left, Inches(4), width, height)
        tf = date_box.text_frame
        if tf is not None:
            tf.text = datetime.now().strftime('%Y-%m-%d')
            if tf.paragraphs and len(tf.paragraphs) > 0:
                tf.paragraphs[0].font.size = Pt(16)
                tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    def _create_directory_slide(self, prs: Presentation, slide_info: Dict, instruction: Dict):
        """
        创建目录幻灯片
        
        Args:
            prs: Presentation对象
            slide_info: 幻灯片信息
            instruction: 完整指令
        """
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        # 添加标题
        title = slide.shapes.title
        if title is not None:
            title.text = "课程目录"
            title.text_frame.paragraphs[0].font.size = Pt(36)
        else:
            # 没有标题占位符，手动添加文本框作为标题
            left = Inches(0.5)
            top = Inches(0.3)
            width = Inches(9)
            height = Inches(1)
            txBox = slide.shapes.add_textbox(left, top, width, height)
            tf = txBox.text_frame
            tf.text = "课程目录"
            tf.paragraphs[0].font.size = Pt(36)
        
        # 添加目录内容
        left = Inches(1.5)
        top = Inches(1.5)
        width = Inches(7)
        height = Inches(3.5)
        
        content_box = slide.shapes.add_textbox(left, top, width, height)
        tf = content_box.text_frame
        if tf is not None:
            tf.word_wrap = True
            
            # 根据教学流程生成目录
            teaching_flow = instruction.get('teaching_flow', [])
            for i, flow in enumerate(teaching_flow):
                if i > 0:
                    p = tf.add_paragraph()
                else:
                    p = tf.paragraphs[0]
                
                if p is not None:
                    stage = flow.get('stage', f'阶段{i+1}')
                    content = flow.get('content', '')
                    p.text = f"{i+1}. {stage} - {content}"
                    p.font.size = Pt(20)
                    p.level = 0
                    p.space_after = Pt(12)
    
    def _create_knowledge_slide(self, prs: Presentation, slide_info: Dict, instruction: Dict):
        """
        创建知识正文幻灯片
        
        Args:
            prs: Presentation对象
            slide_info: 幻灯片信息
            instruction: 完整指令
        """
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        # 添加标题
        title = slide.shapes.title
        if title is not None:
            title.text = slide_info.get('title', '知识点')
            title.text_frame.paragraphs[0].font.size = Pt(32)
        else:
            # 没有标题占位符，手动添加文本框作为标题
            left = Inches(0.5)
            top = Inches(0.2)
            width = Inches(9)
            height = Inches(0.8)
            txBox = slide.shapes.add_textbox(left, top, width, height)
            tf = txBox.text_frame
            tf.text = slide_info.get('title', '知识点')
            tf.paragraphs[0].font.size = Pt(32)
        
        # 添加内容
        left = Inches(0.5)
        top = Inches(1.5)
        width = Inches(9)
        height = Inches(3.5)
        
        content_box = slide.shapes.add_textbox(left, top, width, height)
        tf = content_box.text_frame
        if tf is not None:
            tf.word_wrap = True
            
            content = slide_info.get('content', '')
            if isinstance(content, list):
                content = '\n'.join(content)
            
            # 分段显示内容
            lines = content.split('\n')
            for i, line in enumerate(lines):
                if i > 0:
                    p = tf.add_paragraph()
                else:
                    p = tf.paragraphs[0]
                
                if p is not None:
                    p.text = f"• {line}" if not line.startswith('•') else line
                    p.font.size = Pt(18)
                    p.space_after = Pt(8)
    
    def _create_case_slide(self, prs: Presentation, slide_info: Dict, instruction: Dict):
        """
        创建案例幻灯片
        
        Args:
            prs: Presentation对象
            slide_info: 幻灯片信息
            instruction: 完整指令
        """
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        # 添加标题
        title = slide.shapes.title
        if title is not None:
            title.text = slide_info.get('title', '案例分析')
            title.text_frame.paragraphs[0].font.size = Pt(32)
        else:
            # 没有标题占位符，手动添加文本框作为标题
            left = Inches(0.5)
            top = Inches(0.2)
            width = Inches(9)
            height = Inches(0.8)
            txBox = slide.shapes.add_textbox(left, top, width, height)
            tf = txBox.text_frame
            tf.text = slide_info.get('title', '案例分析')
            tf.paragraphs[0].font.size = Pt(32)
        
        # 添加案例内容
        left = Inches(0.5)
        top = Inches(1.5)
        width = Inches(9)
        height = Inches(3.5)
        
        content_box = slide.shapes.add_textbox(left, top, width, height)
        tf = content_box.text_frame
        if tf is not None:
            tf.word_wrap = True
            
            content = slide_info.get('content', '案例内容待补充')
            tf.text = content
            if tf.paragraphs and len(tf.paragraphs) > 0:
                tf.paragraphs[0].font.size = Pt(18)
        
        # 添加思考问题框
        question_box = slide.shapes.add_textbox(left, Inches(3.2), width, Inches(1))
        qtf = question_box.text_frame
        if qtf is not None:
            qtf.text = "思考问题："
            if qtf.paragraphs and len(qtf.paragraphs) > 0:
                qtf.paragraphs[0].font.size = Pt(16)
                qtf.paragraphs[0].font.bold = True
    
    def _create_interaction_slide(self, prs: Presentation, slide_info: Dict, instruction: Dict):
        """
        创建互动幻灯片
        
        Args:
            prs: Presentation对象
            slide_info: 幻灯片信息
            instruction: 完整指令
        """
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        # 添加标题
        title = slide.shapes.title
        if title is not None:
            title.text = slide_info.get('title', '课堂互动')
            title.text_frame.paragraphs[0].font.size = Pt(32)
        else:
            # 没有标题占位符，手动添加文本框作为标题
            left = Inches(0.5)
            top = Inches(0.2)
            width = Inches(9)
            height = Inches(0.8)
            txBox = slide.shapes.add_textbox(left, top, width, height)
            tf = txBox.text_frame
            tf.text = slide_info.get('title', '课堂互动')
            tf.paragraphs[0].font.size = Pt(32)
        
        # 添加互动内容
        left = Inches(0.5)
        top = Inches(1.5)
        width = Inches(9)
        height = Inches(3)
        
        content_box = slide.shapes.add_textbox(left, top, width, height)
        tf = content_box.text_frame
        if tf is not None:
            tf.word_wrap = True
            
            # 获取课堂活动内容
            activities = instruction.get('classroom_activities', [])
            if activities and isinstance(activities[0], dict):
                # 对象列表格式
                content = '\n'.join([f"• {act.get('activity', '')}: {act.get('description', '')}" for act in activities])
            else:
                # 字符串列表格式
                content = '\n'.join([f"• {act}" for act in activities]) if activities else slide_info.get('content', '互动活动内容')
            
            tf.text = content
            if tf.paragraphs and len(tf.paragraphs) > 0:
                tf.paragraphs[0].font.size = Pt(18)
        
        # 添加互动提示
        hint_box = slide.shapes.add_textbox(left, Inches(3.2), width, Inches(0.8))
        htf = hint_box.text_frame
        if htf is not None:
            htf.text = "请同学们积极参与讨论/回答问题"
            if htf.paragraphs and len(htf.paragraphs) > 0:
                htf.paragraphs[0].font.size = Pt(16)
    
    def _create_summary_slide(self, prs: Presentation, slide_info: Dict, instruction: Dict):
        """
        创建总结幻灯片
        
        Args:
            prs: Presentation对象
            slide_info: 幻灯片信息
            instruction: 完整指令
        """
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        # 添加标题
        title = slide.shapes.title
        if title is not None:
            title.text = "本课重点总结"
            title.text_frame.paragraphs[0].font.size = Pt(32)
        else:
            # 没有标题占位符，手动添加文本框作为标题
            left = Inches(0.5)
            top = Inches(0.2)
            width = Inches(9)
            height = Inches(0.8)
            txBox = slide.shapes.add_textbox(left, top, width, height)
            tf = txBox.text_frame
            tf.text = "本课重点总结"
            tf.paragraphs[0].font.size = Pt(32)
        
        # 添加重点内容
        left = Inches(0.5)
        top = Inches(1.5)
        width = Inches(9)
        height = Inches(3.5)
        
        content_box = slide.shapes.add_textbox(left, top, width, height)
        tf = content_box.text_frame
        if tf is not None:
            tf.word_wrap = True
            
            # 添加重点
            key_points = instruction.get('key_points', [])
            for i, point in enumerate(key_points):
                if i > 0:
                    p = tf.add_paragraph()
                else:
                    p = tf.paragraphs[0]
                
                if p is not None:
                    p.text = f"★ {point}"
                    p.font.size = Pt(20)
                    p.space_after = Pt(10)
            
            # 添加难点
            if tf.paragraphs and len(tf.paragraphs) > 0:
                tf.add_paragraph()
            
            difficult_points = instruction.get('difficult_points', [])
            for i, point in enumerate(difficult_points):
                p = tf.add_paragraph()
                if p is not None:
                    p.text = f"▲ {point}"
                    p.font.size = Pt(18)
                    p.space_after = Pt(8)
    
    def _create_ending_slide(self, prs: Presentation, slide_info: Dict, instruction: Dict):
        """
        创建结尾幻灯片
        
        Args:
            prs: Presentation对象
            slide_info: 幻灯片信息
            instruction: 完整指令
        """
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        # 添加感谢语
        left = Inches(2)
        top = Inches(2)
        width = Inches(6)
        height = Inches(1.5)
        
        thank_box = slide.shapes.add_textbox(left, top, width, height)
        tf = thank_box.text_frame
        if tf is not None:
            tf.text = "感谢聆听"
            if tf.paragraphs and len(tf.paragraphs) > 0:
                tf.paragraphs[0].font.size = Pt(48)
                tf.paragraphs[0].font.bold = True
                tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        # 添加作业提示
        homework = instruction.get('homework', '完成课后练习')
        homework_box = slide.shapes.add_textbox(Inches(1), Inches(3.5), Inches(8), Inches(1))
        htf = homework_box.text_frame
        if htf is not None:
            htf.text = f"课后作业：{homework}"
            if htf.paragraphs and len(htf.paragraphs) > 0:
                htf.paragraphs[0].font.size = Pt(20)
                htf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    def _create_tcp_handshake_single_slide(self, prs: Presentation, instruction: Dict):
        """
        创建TCP三次握手单页完整教学幻灯片（左右分栏布局）
        
        Args:
            prs: Presentation对象
            instruction: 完整指令
        """
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        # 设置淡蓝色渐变背景
        background = slide.background
        fill = background.fill
        fill.gradient()
        fill.gradient_angle = 0
        fill.gradient_stops[0].color.rgb = RGBColor(135, 206, 250)  # 淡蓝色
        fill.gradient_stops[1].color.rgb = RGBColor(240, 248, 255)  # 更淡的蓝色
        
        # 顶部标题区 - 深蓝色横幅
        header_box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            Inches(0), Inches(0), Inches(10), Inches(1.2)
        )
        header_fill = header_box.fill
        header_fill.solid()
        header_fill.fore_color.rgb = RGBColor(0, 102, 204)  # 深蓝色
        
        # 主标题
        title_box = slide.shapes.add_textbox(Inches(0.2), Inches(0.05), Inches(9.6), Inches(0.6))
        tf = title_box.text_frame
        if tf is not None:
            tf.text = instruction.get('course_title', 'TCP三次握手：连接建立的核心机制')
            if tf.paragraphs and len(tf.paragraphs) > 0:
                tf.paragraphs[0].font.size = Pt(36)
                tf.paragraphs[0].font.bold = True
                tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
                tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        # 副标题条
        subtitle_box = slide.shapes.add_textbox(Inches(0.2), Inches(0.7), Inches(9.6), Inches(0.4))
        stf = subtitle_box.text_frame
        if stf is not None:
            stf.text = "理解连接建立过程、seq与ack规则，以及为什么需要三次握手"
            if stf.paragraphs and len(stf.paragraphs) > 0:
                stf.paragraphs[0].font.size = Pt(16)
                stf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
                stf.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        # 左右分栏分隔线
        divider = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(5.8), Inches(1.3), Inches(0.05), Inches(4.1)
        )
        divider_fill = divider.fill
        divider_fill.solid()
        divider_fill.fore_color.rgb = RGBColor(0, 102, 204)
        
        # 左侧：知识点讲解区域
        left_content_title = slide.shapes.add_textbox(Inches(0.2), Inches(1.4), Inches(5.5), Inches(0.4))
        left_title_tf = left_content_title.text_frame
        if left_title_tf is not None:
            left_title_tf.text = "知识点讲解"
            if left_title_tf.paragraphs and len(left_title_tf.paragraphs) > 0:
                left_title_tf.paragraphs[0].font.size = Pt(18)
                left_title_tf.paragraphs[0].font.bold = True
                left_title_tf.paragraphs[0].font.color.rgb = RGBColor(0, 102, 204)
        
        # 左侧内容 - 三次握手流程
        flow_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.2), Inches(1.9), Inches(5.5), Inches(1.5)
        )
        flow_fill = flow_box.fill
        flow_fill.solid()
        flow_fill.fore_color.rgb = RGBColor(255, 255, 255)
        flow_box.line.color.rgb = RGBColor(135, 206, 250)
        flow_box.line.width = Pt(2)
        
        flow_title = slide.shapes.add_textbox(Inches(0.3), Inches(2.0), Inches(5.3), Inches(0.3))
        flow_tf = flow_title.text_frame
        if flow_tf is not None:
            flow_tf.text = "三次握手流程"
            if flow_tf.paragraphs and len(flow_tf.paragraphs) > 0:
                flow_tf.paragraphs[0].font.size = Pt(14)
                flow_tf.paragraphs[0].font.bold = True
                flow_tf.paragraphs[0].font.color.rgb = RGBColor(0, 102, 204)
        
        flow_content = slide.shapes.add_textbox(Inches(0.3), Inches(2.4), Inches(5.3), Inches(0.9))
        flow_content_tf = flow_content.text_frame
        if flow_content_tf is not None:
            flow_content_tf.word_wrap = True
            flow_text = "1. 客户端发送SYN，请求建立连接\n2. 服务端回复SYN+ACK，确认并请求\n3. 客户端回复ACK，确认连接建立"
            flow_content_tf.text = flow_text
            if flow_content_tf.paragraphs and len(flow_content_tf.paragraphs) > 0:
                flow_content_tf.paragraphs[0].font.size = Pt(12)
                # 简单高亮：对整个段落设置格式
                # 注意：python-pptx对文本段内部分字符格式化支持有限，这里简化处理
        
        # 左侧内容 - seq/ack公式
        formula_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.2), Inches(3.5), Inches(5.5), Inches(1.0)
        )
        formula_fill = formula_box.fill
        formula_fill.solid()
        formula_fill.fore_color.rgb = RGBColor(255, 255, 255)
        formula_box.line.color.rgb = RGBColor(135, 206, 250)
        formula_box.line.width = Pt(2)
        
        formula_title = slide.shapes.add_textbox(Inches(0.3), Inches(3.6), Inches(5.3), Inches(0.3))
        formula_tf = formula_title.text_frame
        if formula_tf is not None:
            formula_tf.text = "seq与ack计算规则"
            if formula_tf.paragraphs and len(formula_tf.paragraphs) > 0:
                formula_tf.paragraphs[0].font.size = Pt(14)
                formula_tf.paragraphs[0].font.bold = True
                formula_tf.paragraphs[0].font.color.rgb = RGBColor(0, 102, 204)
        
        formula_content = slide.shapes.add_textbox(Inches(0.3), Inches(4.0), Inches(5.3), Inches(0.4))
        formula_content_tf = formula_content.text_frame
        if formula_content_tf is not None:
            formula_content_tf.word_wrap = True
            formula_text = "seq：发送方序号 | ack = 对方seq + 1"
            formula_content_tf.text = formula_text
            if formula_content_tf.paragraphs and len(formula_content_tf.paragraphs) > 0:
                formula_content_tf.paragraphs[0].font.size = Pt(12)
        
        # 右侧：互动答题模块
        right_title = slide.shapes.add_textbox(Inches(6.0), Inches(1.4), Inches(3.8), Inches(0.4))
        right_title_tf = right_title.text_frame
        if right_title_tf is not None:
            right_title_tf.text = "TCP握手序号大挑战"
            if right_title_tf.paragraphs and len(right_title_tf.paragraphs) > 0:
                right_title_tf.paragraphs[0].font.size = Pt(18)
                right_title_tf.paragraphs[0].font.bold = True
                right_title_tf.paragraphs[0].font.color.rgb = RGBColor(0, 102, 204)
                right_title_tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        # 右侧提示：互动游戏为HTML文件
        game_note = slide.shapes.add_textbox(Inches(6.0), Inches(1.9), Inches(3.8), Inches(0.3))
        game_note_tf = game_note.text_frame
        if game_note_tf is not None:
            game_note_tf.text = "互动答题游戏已生成独立HTML文件"
            if game_note_tf.paragraphs and len(game_note_tf.paragraphs) > 0:
                game_note_tf.paragraphs[0].font.size = Pt(10)
                game_note_tf.paragraphs[0].font.color.rgb = RGBColor(128, 128, 128)
                game_note_tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        # 右侧互动区域背景
        quiz_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(6.0), Inches(2.3), Inches(3.8), Inches(2.9)
        )
        quiz_fill = quiz_box.fill
        quiz_fill.solid()
        quiz_fill.fore_color.rgb = RGBColor(255, 248, 220)  # 浅黄色背景
        quiz_box.line.color.rgb = RGBColor(255, 165, 0)
        quiz_box.line.width = Pt(3)
        
        # 题目1
        q1_box = slide.shapes.add_textbox(Inches(6.2), Inches(2.5), Inches(3.4), Inches(0.6))
        q1_tf = q1_box.text_frame
        if q1_tf is not None:
            q1_tf.word_wrap = True
            q1_tf.text = "第1题：TCP第一次握手发送什么报文？"
            if q1_tf.paragraphs and len(q1_tf.paragraphs) > 0:
                q1_tf.paragraphs[0].font.size = Pt(11)
                q1_tf.paragraphs[0].font.bold = True
        
        q1_options = slide.shapes.add_textbox(Inches(6.2), Inches(3.2), Inches(3.4), Inches(0.5))
        q1_options_tf = q1_options.text_frame
        if q1_options_tf is not None:
            q1_options_tf.word_wrap = True
            q1_options_tf.text = "A. SYN  B. SYN+ACK  C. ACK  D. FIN"
            if q1_options_tf.paragraphs and len(q1_options_tf.paragraphs) > 0:
                q1_options_tf.paragraphs[0].font.size = Pt(10)
        
        q1_answer = slide.shapes.add_textbox(Inches(6.2), Inches(3.8), Inches(3.4), Inches(0.3))
        q1_answer_tf = q1_answer.text_frame
        if q1_answer_tf is not None:
            q1_answer_tf.text = "答案：A"
            if q1_answer_tf.paragraphs and len(q1_answer_tf.paragraphs) > 0:
                q1_answer_tf.paragraphs[0].font.size = Pt(10)
                q1_answer_tf.paragraphs[0].font.color.rgb = RGBColor(0, 128, 0)
                q1_answer_tf.paragraphs[0].font.bold = True
        
        # 题目2
        q2_box = slide.shapes.add_textbox(Inches(6.2), Inches(4.0), Inches(3.4), Inches(0.6))
        q2_tf = q2_box.text_frame
        if q2_tf is not None:
            q2_tf.word_wrap = True
            q2_tf.text = "第2题：客户端seq=100，服务端ack是多少？"
            if q2_tf.paragraphs and len(q2_tf.paragraphs) > 0:
                q2_tf.paragraphs[0].font.size = Pt(11)
                q2_tf.paragraphs[0].font.bold = True
        
        q2_options = slide.shapes.add_textbox(Inches(6.2), Inches(4.7), Inches(3.4), Inches(0.5))
        q2_options_tf = q2_options.text_frame
        if q2_options_tf is not None:
            q2_options_tf.word_wrap = True
            q2_options_tf.text = "A. 100  B. 101  C. 99  D. 102"
            if q2_options_tf.paragraphs and len(q2_options_tf.paragraphs) > 0:
                q2_options_tf.paragraphs[0].font.size = Pt(10)
        
        q2_answer = slide.shapes.add_textbox(Inches(6.2), Inches(5.05), Inches(3.4), Inches(0.3))
        q2_answer_tf = q2_answer.text_frame
        if q2_answer_tf is not None:
            q2_answer_tf.text = "答案：B"
            if q2_answer_tf.paragraphs and len(q2_answer_tf.paragraphs) > 0:
                q2_answer_tf.paragraphs[0].font.size = Pt(10)
                q2_answer_tf.paragraphs[0].font.color.rgb = RGBColor(0, 128, 0)
                q2_answer_tf.paragraphs[0].font.bold = True
        
        # 底部总结区 - 深蓝色条
        footer_box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(5.2), Inches(10), Inches(0.425)
        )
        footer_fill = footer_box.fill
        footer_fill.solid()
        footer_fill.fore_color.rgb = RGBColor(0, 102, 204)
        
        # 总结文本
        summary_box = slide.shapes.add_textbox(Inches(0.2), Inches(5.25), Inches(9.6), Inches(0.35))
        sum_tf = summary_box.text_frame
        if sum_tf is not None:
            sum_tf.text = "三次握手的本质：通过双方确认，建立可靠连接，同时协调初始序号。"
            if sum_tf.paragraphs and len(sum_tf.paragraphs) > 0:
                sum_tf.paragraphs[0].font.size = Pt(14)
                sum_tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
                sum_tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    def _create_tcp_handshake_slide(self, prs: Presentation, slide_info: Dict, instruction: Dict):
        """
        创建TCP三次握手专用幻灯片（用于多页模式）
        
        Args:
            prs: Presentation对象
            slide_info: 幻灯片信息
            instruction: 完整指令
        """
        # 使用标题作为幻灯片标题
        slide_title = slide_info.get('title', 'TCP三次握手：连接建立的核心机制')
        enhanced_instruction = instruction.copy()
        enhanced_instruction['course_title'] = slide_title
        
        # 调用单页方法，使用左右分栏布局
        self._create_tcp_handshake_single_slide(prs, enhanced_instruction)
    
    def _get_default_structure(self, instruction: Dict) -> List[Dict]:
        """
        获取默认PPT结构
        
        Args:
            instruction: 融合指令
            
        Returns:
            默认幻灯片结构列表
        """
        course_title = instruction.get('course_title', '教学课件')
        knowledge_points = instruction.get('knowledge_points', [])
        
        structure = [
            {'slide_type': 'cover', 'title': course_title},
            {'slide_type': 'directory', 'title': '课程目录'},
        ]
        
        # 添加知识点幻灯片
        for i, point in enumerate(knowledge_points):
            structure.append({
                'slide_type': 'knowledge',
                'title': f'知识点{i+1}',
                'content': point
            })
        
        # 添加案例和互动
        structure.append({'slide_type': 'case', 'title': '案例分析', 'content': '相关案例'})
        structure.append({'slide_type': 'interaction', 'title': '课堂互动', 'content': '互动活动'})
        structure.append({'slide_type': 'summary', 'title': '重点总结'})
        structure.append({'slide_type': 'ending', 'title': '课程结束'})
        
        return structure


# 全局PPT生成器实例
_ppt_generator = None


def get_ppt_generator() -> PPTGenerator:
    """
    获取全局PPT生成器实例（单例模式）
    
    Returns:
        PPTGenerator实例
    """
    global _ppt_generator
    if _ppt_generator is None:
        _ppt_generator = PPTGenerator()
    return _ppt_generator
